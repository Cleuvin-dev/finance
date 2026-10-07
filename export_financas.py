import openpyxl, warnings, json, datetime, pathlib
warnings.filterwarnings('ignore')
w=openpyxl.load_workbook('FINANÇAS.xlsx',data_only=False)
d=openpyxl.load_workbook('FINANÇAS.xlsx',data_only=True)
names=w.sheetnames[8:20]; records=[]
def val(s,a): return d[s][a].value
def date(v): return v.date().isoformat() if isinstance(v,datetime.datetime) else ''
for m,s in enumerate(names,1):
 for r in range(13,53):
  for kind,dc,cc,dt,vc,sc,bc in [('fixed','B','C','D','F','G','H'),('variable','J','K','L','N','O','P')]:
   desc=val(s,f'{dc}{r}'); amount=val(s,f'{vc}{r}')
   if desc and isinstance(amount,(int,float)):
    rec={'id':f'{m}-{kind}-{r}','month':m,'kind':kind,'description':str(desc),'category':str(val(s,f'{cc}{r}') or ''),'date':date(val(s,f'{dt}{r}')),'cents':round(amount*100),'paid':val(s,f'{sc}{r}')=='PAGO','bank':str(val(s,f'{bc}{r}') or ''),'source':f'{s}!{vc}{r}'}
    if w[s][f'{vc}{r}'].data_type=='f' and m>1: rec['recurringFrom']=f'{m-1}-{kind}-{r}'
    records.append(rec)
 for kind,rows,dc,dt,vc,sc in [('card',range(63,70),'B',None,'D','C'),('income',range(63,70),'F','G','H',None),('investment',range(15,44),'V','W','X',None)]:
  for r in rows:
   amount=val(s,f'{vc}{r}')
   if isinstance(amount,(int,float)) and amount:
    records.append({'id':f'{m}-{kind}-{r}','month':m,'kind':kind,'description':str(val(s,f'{dc}{r}') or kind),'category':'Cartão de crédito' if kind=='card' else '', 'date':date(val(s,f'{dt}{r}')) if dt else '', 'cents':round(amount*100),'paid':val(s,f'{sc}{r}')=='PAGO' if sc else True,'bank':str(val(s,f'{dc}{r}') or '') if kind=='card' else '', 'source':f'{s}!{vc}{r}'})
reserve=d[w.sheetnames[5]]
for m in range(1,13):
 amount=reserve.cell(9,m+6).value or 0
 if amount: records.append({'id':f'{m}-reserve','month':m,'kind':'reserve','description':'Reserva de emergência','category':'Reserva','date':f'2026-{m:02d}-01','dateApproximate':True,'cents':round(amount*100),'paid':True,'bank':'','source':f'{reserve.title}!{reserve.cell(9,m+6).coordinate}'})
categories=list(dict.fromkeys(str(d['Categorias'].cell(r,2).value).strip().replace('Eduação','Educação') for r in range(6,15)))
for rec in records: rec['category']=rec['category'].strip().replace('Eduação','Educação')
data={'year':2026,'records':records,'categories':categories,'banks':[d['Bancos'].cell(r,2).value for r in range(6,13)],'cards':[d[w.sheetnames[7]].cell(r,2).value for r in range(6,12)],'monthlyTargets':[round((reserve.cell(8,m+6).value or 0)*100) for m in range(1,13)],'allocation':[50,30,20],'essentialMonthlyCents':189000,'wealthCents':0,'wealthGoals':[5000000,10000000,15000000,20000000,25000000],'monthStates':['realizado']*8+['previsto']*4,'sourceSheets':w.sheetnames,'help':[str(c.value) for row in w['Ajuda'] for c in row if c.value and c.data_type!='f']}
out=pathlib.Path('controle-financeiro/lib'); out.mkdir(exist_ok=True)
(out/'initial-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
print('Imported',len(records),'records;',len(data['help']),'help cells')
