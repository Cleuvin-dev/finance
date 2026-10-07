import openpyxl, warnings, sys, collections, json
sys.stdout.reconfigure(encoding='utf-8')
warnings.filterwarnings('ignore')
w=openpyxl.load_workbook('FINANÇAS.xlsx',data_only=False)
d=openpyxl.load_workbook('FINANÇAS.xlsx',data_only=True)
months=w.sheetnames[8:20]
errors=[]; broken=[]
for s in w:
 for row in s:
  for c in row:
   if d[s.title][c.coordinate].data_type=='e': errors.append((s.title,c.coordinate,d[s.title][c.coordinate].value))
   if c.data_type=='f' and '#REF!' in str(c.value): broken.append((s.title,c.coordinate,c.value))
print('ERRORS',errors,'BROKEN',broken)
cats=collections.Counter(); desc=collections.Counter(); mismatches=[]
for name in months:
 s=d[name]; f=w[name]
 totals={a:s[a].value for a in ['D5','F5','H5','J5','L5','N5','D10','M10','D59','D60','S15']}
 print('MONTH',name,totals)
 detail=0; paid=0
 for r in range(13,53):
  for dc,cc,vc,sc,st in [('B','C','F','G','E'),('J','K','N','O','M')]:
   v=s[f'{vc}{r}'].value
   if isinstance(v,(int,float)) and v:
    detail+=v
    if s[f'{sc}{r}'].value=='PAGO': paid+=v
    if name in months[:8]:
     cats[str(s[f'{cc}{r}'].value)]+=v; desc[str(s[f'{dc}{r}'].value)]+=v
    if s[f'{sc}{r}'].value=='PAGO' and s[f'{st}{r}'].value!='PAGO': mismatches.append((name,f'{st}{r}',s[f'{st}{r}'].value))
 print('RECON',detail,paid,'cards',s['D59'].value,s['D60'].value)
 if name in ['Janeiro','Agosto','Setembro']:
  for r in range(54,71):
   vals=[(c.coordinate,str(c.value),str(d[name][c.coordinate].value)) for c in f[r] if c.value is not None]
   if vals: print('EXTRA',name,vals)
print('CATEGORIES',cats,'TOP_DESCRIPTIONS',desc.most_common(8),'PAID_STATUS_MISMATCH',len(mismatches),mismatches[:6])
for name in ['Rendimento','Ajuda']:
 print('OTHER',name)
 for row in w[name]:
  vals=[(c.coordinate,str(c.value)[:250],str(d[name][c.coordinate].value)) for c in row if c.value is not None]
  if vals: print(vals)
  if row[0].row>45: break
for name in ['DASHBOARD','DM','Tabela dinâmica 2']:
 for ch in w[name]._charts:
  print('CHART',name,[(str(se.val),str(se.cat)) for se in ch.series])
print('NAMES',[(k,str(v)) for k,v in w.defined_names.items()])
